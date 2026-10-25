import openpyxl

wb = openpyxl.load_workbook('COUGAR EXE EQ.xlsx', data_only=True)
sheet = wb['Bikes']
products = []
for r in range(8, sheet.max_row + 1):
    sr = sheet.cell(r, 1).value
    name = sheet.cell(r, 2).value
    desc = sheet.cell(r, 3).value
    qty = sheet.cell(r, 4).value
    cat = sheet.cell(r, 5).value
    color = sheet.cell(r, 6).value
    price = sheet.cell(r, 7).value
    discount = sheet.cell(r, 9).value
    if name is not None:
        products.append({
            'row': r,
            'sr': sr,
            'name': name,
            'price': price,
        })

print(f'Total in COUGAR: {len(products)}')
for p in products:
    print(f"Row {p['row']:3d} | Sr: {p['sr']} | {p['name']} | Price: {p['price']}")
